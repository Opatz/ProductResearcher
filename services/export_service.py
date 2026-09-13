import logging
import json
import re
from pathlib import Path
from typing import Any, List, Dict, Union, Optional
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

logger = logging.getLogger(__name__)


def _format_list_field(val: Any) -> str:
    """Formatiert Listen oder Arrays in einen sauberen, lesbaren String (Semikolon-separiert)."""
    if isinstance(val, (list, tuple)):
        clean_items = [str(x).strip() for x in val if x is not None and str(x).strip()]
        return "; ".join(clean_items)
    elif val is None:
        return ""
    return str(val).strip()


class ExportService:
    """Service zum Konvertieren und Exportieren von Pipeline-Ergebnissen in CSV und 2-Sheet Excel (.xlsx)."""

    def __init__(self, output_dir: Path):
        self.output_dir = Path(output_dir)
        self.output_dir.mkdir(parents=True, exist_ok=True)

    @staticmethod
    def flatten_record(record: Dict[str, Any]) -> Dict[str, Any]:
        """Flacht verschachtelte Felder oder Listen in lesbare Tabellenspalten ab."""
        flat_row = {}
        for k, v in record.items():
            if isinstance(v, (list, tuple)):
                flat_row[k] = _format_list_field(v)
            elif isinstance(v, dict):
                for sub_k, sub_v in v.items():
                    if isinstance(sub_v, (list, tuple)):
                        flat_row[f"{k}_{sub_k}"] = _format_list_field(sub_v)
                    elif isinstance(sub_v, dict):
                        flat_row[f"{k}_{sub_k}"] = json.dumps(sub_v, ensure_ascii=False)
                    else:
                        flat_row[f"{k}_{sub_k}"] = sub_v
            else:
                flat_row[k] = v
        return flat_row

    @classmethod
    def build_web_research_rows(
        cls,
        item_id: str,
        listings: List[Union[Dict[str, Any], Any]]
    ) -> List[Dict[str, Any]]:
        """
        Erstellt strukturierte Zeilen für das Sheet 'Marktrecherche_Webpreise' mit klickbaren
        =HYPERLINK(...)-Formeln und Standardspalten.
        """
        rows = []
        for listing in (listings or []):
            if hasattr(listing, "model_dump"):
                listing_dict = listing.model_dump(mode="json")
            elif isinstance(listing, dict):
                listing_dict = listing
            else:
                continue

            site_name = str(listing_dict.get("website_name") or "").strip()
            title = str(listing_dict.get("listing_titel") or "").strip()
            preis = listing_dict.get("preis_eur")
            preis_typ = listing_dict.get("preis_typ")
            if hasattr(preis_typ, "value"):
                preis_typ = preis_typ.value
            match_gen = listing_dict.get("match_genauigkeit")
            if hasattr(match_gen, "value"):
                match_gen = match_gen.value
            zustand = str(listing_dict.get("zustand_referenz") or "").strip()
            status = listing_dict.get("recherche_status") or listing_dict.get("status")
            if hasattr(status, "value"):
                status = status.value

            url = str(listing_dict.get("quell_url") or listing_dict.get("target_url") or "").strip()

            if url.startswith("http://") or url.startswith("https://"):
                friendly_text = title if title else (f"Link ({site_name})" if site_name else "Link öffnen")
                safe_url = url.replace('"', '""')
                safe_friendly = friendly_text.replace('"', '""')
                link_formula = f'=HYPERLINK("{safe_url}", "{safe_friendly}")'
            else:
                link_formula = url or "N/A"

            rows.append({
                "Artikel_ID": str(item_id or "N/A").strip(),
                "Website_Name": site_name,
                "Produkt_Titel": title,
                "Preis_EUR": preis,
                "Preis_Typ": str(preis_typ or "").strip(),
                "Match_Genauigkeit": str(match_gen or "").strip(),
                "Zustand_Referenz": zustand,
                "Status": str(status or "").strip(),
                "Link_URL": link_formula
            })

        return rows

    @classmethod
    def build_unified_item_row_from_state(cls, state: Any) -> Dict[str, Any]:
        """Extrahiert alle Felder aus einem PipelineState-Objekt."""
        return cls.build_unified_item_row(
            visual_analysis=getattr(state, "visual_analysis_json", None),
            synthesis=getattr(state, "retail_price_synthesis_json", None) or getattr(state, "web_price_summary", None),
            reference_listings=getattr(state, "reference_listings", None) or getattr(state, "discovered_web_sources", None),
            detected_id=getattr(state, "detected_id", None),
            item_name=getattr(state, "item_name", ""),
            video_filename=state.video_path.name if getattr(state, "video_path", None) else "",
            images_count=len(getattr(state, "image_paths", []))
        )

    @classmethod
    def build_unified_item_row(
        cls,
        visual_analysis: Optional[Dict[str, Any]] = None,
        synthesis: Optional[Union[Dict[str, Any], Any]] = None,
        reference_listings: Optional[List[Any]] = None,
        detected_id: Optional[str] = None,
        item_name: str = "",
        video_filename: str = "",
        images_count: int = 0,
        **kwargs: Any
    ) -> Dict[str, Any]:
        """
        Führt sämtliche extrahierten Details der 3-Stufen-Marktbewertung in eine einheitliche,
        vollständige Tabellenzeile für die Hauptempfehlung zusammen.
        """
        vis_dict = visual_analysis if isinstance(visual_analysis, dict) else {}
        
        synth_dict = {}
        if synthesis is not None:
            if hasattr(synthesis, "model_dump"):
                synth_dict = synthesis.model_dump(mode="json")
            elif isinstance(synthesis, dict):
                synth_dict = synthesis

        vis_merkmale = vis_dict.get("physische_merkmale") if isinstance(vis_dict.get("physische_merkmale"), dict) else {}
        vis_zustand = vis_dict.get("zustandsbericht") if isinstance(vis_dict.get("zustandsbericht"), dict) else {}
        synth_merkmale = synth_dict.get("physische_merkmale") if isinstance(synth_dict.get("physische_merkmale"), dict) else {}
        synth_zustand = synth_dict.get("zustandsbericht") if isinstance(synth_dict.get("zustandsbericht"), dict) else {}

        # 1. ID - Strikte Priorität auf die interne Artikel-ID
        folder_match = re.search(r"Artikel_([A-Za-z0-9_]+)", str(item_name))
        folder_id = folder_match.group(1) if folder_match else ""

        if detected_id and str(detected_id).strip().lower() not in {"n/a", "none", "null", ""}:
            item_id = str(detected_id).strip()
        elif folder_id:
            item_id = str(folder_id).strip()
        else:
            item_id = str(item_name or "N/A").strip()

        # 2. Titel & Beschreibung
        titel = synth_dict.get("titel") or vis_dict.get("titel") or item_name
        produktbeschreibung = synth_dict.get("produktbeschreibung") or vis_dict.get("produktbeschreibung") or ""

        # 3. Stammdaten & Attribute
        kategorie = vis_dict.get("kategorie") or ""
        hersteller = vis_dict.get("hersteller_oder_marke") or ""
        epoche = vis_dict.get("modell_oder_epoche") or ""
        geschaetztes_jahr = vis_dict.get("geschaetztes_jahr_oder_epoche") or ""
        authentizitaet = vis_dict.get("authentizitaet") or "unklar"

        raw_stempel = vis_dict.get("erkannte_nummern_oder_stempel")
        stempel = _format_list_field(raw_stempel)

        material = synth_merkmale.get("material") or vis_merkmale.get("material") or ""
        farbe = synth_merkmale.get("farbe") or vis_merkmale.get("farbe") or ""
        zustand = synth_zustand.get("zustand") or vis_zustand.get("zustand") or ""

        raw_maengel = synth_zustand.get("maengel") if synth_zustand.get("maengel") is not None else vis_zustand.get("maengel")
        maengel = _format_list_field(raw_maengel)

        raw_fehlende = synth_zustand.get("fehlende_teile") if synth_zustand.get("fehlende_teile") is not None else vis_zustand.get("fehlende_teile")
        fehlende_teile = _format_list_field(raw_fehlende)

        # 4. Maße & Logistik
        laenge_cm = synth_merkmale.get("laenge_cm") if synth_merkmale.get("laenge_cm") is not None else vis_merkmale.get("laenge_cm")
        breite_cm = synth_merkmale.get("breite_cm") if synth_merkmale.get("breite_cm") is not None else vis_merkmale.get("breite_cm")
        hoehe_cm = synth_merkmale.get("hoehe_cm") if synth_merkmale.get("hoehe_cm") is not None else vis_merkmale.get("hoehe_cm")
        durchmesser_cm = synth_merkmale.get("durchmesser_cm") if synth_merkmale.get("durchmesser_cm") is not None else vis_merkmale.get("durchmesser_cm")
        gewicht_kg = synth_merkmale.get("gewicht_kg") if synth_merkmale.get("gewicht_kg") is not None else vis_merkmale.get("gewicht_kg")
        logistik_kategorie = synth_merkmale.get("logistik_kategorie") or vis_merkmale.get("logistik_kategorie") or ""

        # 5. Appraiser LLM Preissynthese & Rechercheergebnisse
        retail_preis = synth_dict.get("geschaetzter_retail_preis_eur", 0.0)
        preisspanne_min = synth_dict.get("preisspanne_min_eur", 0.0)
        preisspanne_max = synth_dict.get("preisspanne_max_eur", 0.0)
        median_web = synth_dict.get("median_web_preis_eur", 0.0)
        anzahl_preise = synth_dict.get("anzahl_gefundene_preise", 0)
        begruendung_preise = synth_dict.get("begruendung_preisfindung", "")
        ausreisser_notiz = synth_dict.get("ausreisser_bereinigung_notiz", "")

        # Top Referenz Links
        top_links = []
        if reference_listings:
            for rl in reference_listings:
                if hasattr(rl, "model_dump"):
                    rd = rl.model_dump(mode="json")
                elif isinstance(rl, dict):
                    rd = rl
                else:
                    continue
                r_url = str(rd.get("quell_url") or rd.get("target_url") or "").strip()
                r_stat = str(rd.get("recherche_status") or rd.get("status") or "").lower()
                if (r_url.startswith("http://") or r_url.startswith("https://")) and (not r_stat or "erfolg" in r_stat):
                    if r_url not in top_links:
                        top_links.append(r_url)
        top_links_str = "; ".join(top_links[:5])

        return {
            # Stammdaten & Identifikation
            "id": item_id,
            "titel": titel,
            "produktbeschreibung": produktbeschreibung,
            "kategorie": kategorie,
            "hersteller_oder_marke": hersteller,
            "modell_oder_epoche": epoche,
            "geschaetztes_jahr_oder_epoche": geschaetztes_jahr,
            "authentizitaet": authentizitaet,
            "erkannte_nummern_oder_stempel": stempel,
            "material": material,
            "farbe": farbe,
            "zustand": zustand,
            "maengel": maengel,
            "fehlende_teile": fehlende_teile,

            # Maße & Logistik
            "laenge_cm": laenge_cm,
            "breite_cm": breite_cm,
            "hoehe_cm": hoehe_cm,
            "durchmesser_cm": durchmesser_cm,
            "gewicht_kg": gewicht_kg,
            "logistik_kategorie": logistik_kategorie,

            # Marktbewertung (Appraiser LLM Synthese)
            "Empfohlener_Retail_Preis_EUR": retail_preis,
            "Preisspanne_Min_EUR": preisspanne_min,
            "Preisspanne_Max_EUR": preisspanne_max,
            "Median_Web_Preis_EUR": median_web,
            "Anzahl_gefundene_Webpreise": anzahl_preise,
            "Begruendung_Preisfindung": begruendung_preise,
            "Ausreisser_Bereinigung_Notiz": ausreisser_notiz,
            "Top_Referenz_Links": top_links_str,

            # Prozess- & Dateimetadaten
            "ordner_name": item_name,
            "video_datei": video_filename,
            "anzahl_bilder": images_count
        }

    def _apply_excel_styling(self, workbook: openpyxl.Workbook) -> None:
        """Formatiert alle Sheets mit ansprechenden Spaltenbreiten, Kopfzeilen, Link-Formatierung und Textumbruch."""
        header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        data_font = Font(name="Calibri", size=10)
        link_font = Font(name="Calibri", size=10, color="FF0563C1", underline="single")
        thin_border = Border(
            left=Side(style='thin', color='D9D9D9'),
            right=Side(style='thin', color='D9D9D9'),
            top=Side(style='thin', color='D9D9D9'),
            bottom=Side(style='thin', color='D9D9D9')
        )

        for sheetname in workbook.sheetnames:
            ws = workbook[sheetname]
            ws.freeze_panes = "A2"
            ws.row_dimensions[1].height = 28

            # Header formatieren
            for col_idx in range(1, ws.max_column + 1):
                cell = ws.cell(row=1, column=col_idx)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

            # Datenzeilen formatieren & Spaltenbreiten berechnen
            for col_idx in range(1, ws.max_column + 1):
                col_letter = get_column_letter(col_idx)
                max_len = 0

                header_val = ws.cell(row=1, column=col_idx).value
                if header_val:
                    max_len = max(max_len, len(str(header_val)))

                for row_idx in range(2, ws.max_row + 1):
                    cell = ws.cell(row=row_idx, column=col_idx)
                    cell.border = thin_border

                    val = cell.value
                    if val is not None:
                        val_str = str(val)
                        if val_str.startswith("=HYPERLINK("):
                            cell.font = link_font
                            cell.alignment = Alignment(horizontal="left", vertical="top")
                            m = re.search(r'=HYPERLINK\(".*?",\s*"(.*?)"\)', val_str)
                            display_text = m.group(1) if m else val_str
                            max_len = max(max_len, len(display_text))
                        else:
                            cell.font = data_font
                            first_line = val_str.split("\n")[0] if "\n" in val_str else val_str
                            max_len = max(max_len, len(first_line))

                            if isinstance(val, (int, float)):
                                cell.alignment = Alignment(horizontal="right", vertical="top")
                            else:
                                cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
                    else:
                        cell.font = data_font
                        cell.alignment = Alignment(horizontal="left", vertical="top")

                ws.column_dimensions[col_letter].width = max(12, min(max_len + 3, 55))

    def export_single_item(self, state: Any, base_filename: str = "item_result") -> Dict[str, Path]:
        """
        Exportiert einen einzelnen Artikel in CSV und ein sauberes 2-Sheet Excel (.xlsx):
        - Sheet 1: Hauptempfehlung (Alle Attribute & Retail-Bewertung)
        - Sheet 2: Marktrecherche_Webpreise (10 untersuchte Webseiten mit klickbaren =HYPERLINK Formeln)
        """
        main_row = self.build_unified_item_row_from_state(state)
        df_main = pd.DataFrame([main_row])

        tracking_cols = ["VerkaufsOrt", "Status", "ErzielterPreis"]
        for col in tracking_cols:
            df_main[col] = ""
        other_cols = [c for c in df_main.columns if c not in tracking_cols]
        df_main = df_main[other_cols + tracking_cols]

        csv_path = self.output_dir / f"{base_filename}.csv"
        excel_path = self.output_dir / f"{base_filename}.xlsx"

        # CSV-Export
        df_main.to_csv(csv_path, index=False, sep=";", encoding="utf-8-sig")
        logger.info(f"CSV exportiert: {csv_path}")

        # Web-Research Listings für Sheet 2
        ref_listings = (
            getattr(state, "reference_listings", None) 
            or getattr(state, "discovered_web_sources", None)
            or []
        )
        item_id = main_row.get("id") or getattr(state, "detected_id", None) or "N/A"
        web_research_rows = self.build_web_research_rows(item_id, ref_listings)

        # 2-Sheet Excel-Export
        with pd.ExcelWriter(excel_path, engine="openpyxl") as writer:
            # Sheet 1: Hauptempfehlung
            df_main.to_excel(writer, index=False, sheet_name="Hauptempfehlung")

            # Sheet 2: Marktrecherche_Webpreise
            if web_research_rows:
                df_web = pd.DataFrame(web_research_rows)
            else:
                df_web = pd.DataFrame(columns=[
                    "Artikel_ID", "Website_Name", "Produkt_Titel", "Preis_EUR",
                    "Preis_Typ", "Match_Genauigkeit", "Zustand_Referenz", "Status", "Link_URL"
                ])
            df_web.to_excel(writer, index=False, sheet_name="Marktrecherche_Webpreise")

        wb = openpyxl.load_workbook(excel_path)
        self._apply_excel_styling(wb)
        wb.save(excel_path)
        logger.info(f"Excel mit 2 Sheets exportiert: {excel_path}")

        return {
            "csv": csv_path,
            "excel": excel_path,
            "row_count": 1
        }

    def export_consolidated_batch(
        self,
        items_rows: List[Dict[str, Any]],
        measures_rows: Optional[List[Dict[str, Any]]] = None,
        web_research_rows: Optional[List[Dict[str, Any]]] = None,
        base_filename: str = "consolidated_execution_results"
    ) -> Dict[str, Path]:
        """
        Exportiert die konsolidierte Übersicht aller Artikel in CSV und 2-Sheet Excel.
        - Sheet 1: Alle_Artikel
        - Sheet 2: Marktrecherche_Webpreise
        """
        df_items = pd.DataFrame(items_rows)

        tracking_cols = ["VerkaufsOrt", "Status", "ErzielterPreis"]
        for col in tracking_cols:
            if col not in df_items.columns:
                df_items[col] = ""
            else:
                df_items[col] = df_items[col].fillna("")

        other_cols = [c for c in df_items.columns if c not in tracking_cols]
        df_items = df_items[other_cols + tracking_cols]

        csv_path = self.output_dir / f"{base_filename}.csv"
        excel_path = self.output_dir / f"{base_filename}.xlsx"

        # CSV
        try:
            df_items.to_csv(csv_path, index=False, sep=";", encoding="utf-8-sig")
            logger.info(f"Konsolidierte CSV exportiert: {csv_path}")
        except PermissionError:
            csv_path = self.output_dir / f"{base_filename}_neu.csv"
            df_items.to_csv(csv_path, index=False, sep=";", encoding="utf-8-sig")
            logger.warning(f"Zieldatei war gesperrt. Konsolidierte CSV gespeichert unter: {csv_path}")

        # Excel
        def _write_excel(target_path: Path):
            with pd.ExcelWriter(target_path, engine="openpyxl") as writer:
                df_items.to_excel(writer, index=False, sheet_name="Alle_Artikel")

                if web_research_rows is not None:
                    if web_research_rows:
                        df_web = pd.DataFrame(web_research_rows)
                    else:
                        df_web = pd.DataFrame(columns=[
                            "Artikel_ID", "Website_Name", "Produkt_Titel", "Preis_EUR",
                            "Preis_Typ", "Match_Genauigkeit", "Zustand_Referenz", "Status", "Link_URL"
                        ])
                    df_web.to_excel(writer, index=False, sheet_name="Marktrecherche_Webpreise")

            wb = openpyxl.load_workbook(target_path)
            self._apply_excel_styling(wb)
            wb.save(target_path)

        try:
            _write_excel(excel_path)
            logger.info(f"Konsolidierte Excel exportiert: {excel_path}")
        except PermissionError:
            excel_path = self.output_dir / f"{base_filename}_neu.xlsx"
            _write_excel(excel_path)
            logger.warning(f"Zieldatei war in Excel geöffnet. Konsolidierte Excel gespeichert unter: {excel_path}")

        return {
            "csv": csv_path,
            "excel": excel_path,
            "row_count": len(items_rows)
        }

    def export_data(
        self, 
        data: Union[Dict[str, Any], List[Dict[str, Any]]], 
        base_filename: str = "pipeline_result"
    ) -> Dict[str, Path]:
        """
        Kompatibilitäts-Exportfunktion für bestehende Test- und Direktexporte.
        """
        if isinstance(data, list):
            cleaned_records = [self.flatten_record(r) for r in data]
            df_main = pd.DataFrame(cleaned_records)
            web_rows = None
        elif isinstance(data, dict):
            df_main = pd.DataFrame([self.flatten_record(data)])
            ref_listings = data.get("reference_listings") or data.get("discovered_web_sources")
            web_rows = self.build_web_research_rows(data.get("id", "N/A"), ref_listings) if ref_listings else None
        else:
            raise ValueError(f"Ungültiger Datentyp für Export: {type(data)}")

        csv_path = self.output_dir / f"{base_filename}.csv"
        excel_path = self.output_dir / f"{base_filename}.xlsx"

        df_main.to_csv(csv_path, index=False, sep=";", encoding="utf-8-sig")
        logger.info(f"CSV exportiert: {csv_path}")

        with pd.ExcelWriter(excel_path, engine="openpyxl") as writer:
            df_main.to_excel(writer, index=False, sheet_name="Hauptempfehlung")
            if web_rows:
                df_web = pd.DataFrame(web_rows)
                df_web.to_excel(writer, index=False, sheet_name="Marktrecherche_Webpreise")

        wb = openpyxl.load_workbook(excel_path)
        self._apply_excel_styling(wb)
        wb.save(excel_path)
        logger.info(f"Excel exportiert: {excel_path}")

        return {
            "csv": csv_path,
            "excel": excel_path,
            "row_count": len(df_main)
        }
