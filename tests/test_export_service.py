import unittest
import tempfile
import shutil
from pathlib import Path
import openpyxl
import pandas as pd

from pipeline.models import (
    PriceType,
    MatchGenauigkeit,
    RechercheStatus,
    ReferenceListing,
    RetailPriceSynthesis,
)
from pipeline.state import PipelineState
from services.export_service import ExportService


class TestExportService(unittest.TestCase):
    """Testfälle für den 2-Sheet Excel- und CSV-Export gem. ADR 0002 und ADR 0003."""

    def setUp(self):
        self.test_dir = Path(tempfile.mkdtemp())
        self.export_service = ExportService(self.test_dir)

        # 10 Beispiel-Listings wie aus Phase 2
        self.sample_listings = [
            ReferenceListing(
                website_name="eBay",
                listing_titel="Original Teak Tisch 60er Jahre",
                preis_eur=280.0,
                urspruengliche_waehrung="EUR",
                preis_typ=PriceType.REALISIERTER_VERKAUFSPREIS,
                match_genauigkeit=MatchGenauigkeit.EXAKTER_TREFFER,
                zustand_referenz="Gebraucht mit leichten Gebrauchsspuren",
                quell_url="https://www.ebay.de/itm/123456789",
                recherche_status=RechercheStatus.ERFOLGREICH
            ),
            ReferenceListing(
                website_name="Kleinanzeigen",
                listing_titel="Vintage Couchtisch Teakholz",
                preis_eur=250.0,
                urspruengliche_waehrung="EUR",
                preis_typ=PriceType.ANGEBOTSPREIS,
                match_genauigkeit=MatchGenauigkeit.MODELLVARIANTE,
                zustand_referenz="Guter Erhaltungszustand",
                quell_url="https://www.kleinanzeigen.de/s-anzeige/teak-tisch/987654",
                recherche_status=RechercheStatus.ERFOLGREICH
            ),
            ReferenceListing(
                website_name="Pamono",
                listing_titel="Dänischer Teak Couchtisch Galeriepreis",
                preis_eur=750.0,
                urspruengliche_waehrung="EUR",
                preis_typ=PriceType.ANGEBOTSPREIS,
                match_genauigkeit=MatchGenauigkeit.EXAKTER_TREFFER,
                zustand_referenz="Restauriert",
                quell_url="https://www.pamono.de/vintage-teak-table",
                recherche_status=RechercheStatus.ERFOLGREICH
            ),
            ReferenceListing(
                website_name="1stDibs",
                listing_titel="Mid-Century Modern Coffee Table",
                preis_eur=850.0,
                urspruengliche_waehrung="USD",
                preis_typ=PriceType.ANGEBOTSPREIS,
                match_genauigkeit=MatchGenauigkeit.EXAKTER_TREFFER,
                zustand_referenz="Exzellent",
                quell_url="https://www.1stdibs.com/furniture/tables/111222",
                recherche_status=RechercheStatus.ERFOLGREICH
            ),
            ReferenceListing(
                website_name="Catawiki",
                listing_titel="Teak Salontisch Auktionsergebnis",
                preis_eur=310.0,
                urspruengliche_waehrung="EUR",
                preis_typ=PriceType.REALISIERTER_VERKAUFSPREIS,
                match_genauigkeit=MatchGenauigkeit.EXAKTER_TREFFER,
                zustand_referenz="Altersgemäße Patina",
                quell_url="https://www.catawiki.com/l/333444",
                recherche_status=RechercheStatus.ERFOLGREICH
            ),
            ReferenceListing(
                website_name="Etsy",
                listing_titel="Danish Teak Coffee Table 1960s",
                preis_eur=390.0,
                urspruengliche_waehrung="EUR",
                preis_typ=PriceType.ANGEBOTSPREIS,
                match_genauigkeit=MatchGenauigkeit.MODELLVARIANTE,
                zustand_referenz="Sehr gut",
                quell_url="https://www.etsy.com/listing/555666",
                recherche_status=RechercheStatus.ERFOLGREICH
            ),
            ReferenceListing(
                website_name="LiveAuctioneers",
                listing_titel="Mid Century Modern Table",
                preis_eur=290.0,
                urspruengliche_waehrung="USD",
                preis_typ=PriceType.REALISIERTER_VERKAUFSPREIS,
                match_genauigkeit=MatchGenauigkeit.MODELLVARIANTE,
                zustand_referenz="Gebrauchsspuren",
                quell_url="https://www.liveauctioneers.com/item/777888",
                recherche_status=RechercheStatus.ERFOLGREICH
            ),
            ReferenceListing(
                website_name="Dorotheum",
                listing_titel="Katalogsuche kein Treffer",
                preis_eur=None,
                urspruengliche_waehrung="EUR",
                preis_typ=PriceType.UNBEKANNT,
                match_genauigkeit=MatchGenauigkeit.KEIN_TREFFER,
                zustand_referenz="",
                quell_url="https://www.dorotheum.com/lot/none",
                recherche_status=RechercheStatus.KEIN_PREIS_GEFUNDEN
            ),
            ReferenceListing(
                website_name="Quoka",
                listing_titel="Teaktisch",
                preis_eur=200.0,
                urspruengliche_waehrung="EUR",
                preis_typ=PriceType.ANGEBOTSPREIS,
                match_genauigkeit=MatchGenauigkeit.AEHNLICHES_OBJEKT,
                zustand_referenz="Gebraucht",
                quell_url="https://www.quoka.de/999000",
                recherche_status=RechercheStatus.ERFOLGREICH
            ),
            ReferenceListing(
                website_name="AnticStore",
                listing_titel="Table Basse Teck",
                preis_eur=None,
                urspruengliche_waehrung="EUR",
                preis_typ=PriceType.UNBEKANNT,
                match_genauigkeit=MatchGenauigkeit.KEIN_TREFFER,
                zustand_referenz="",
                quell_url="",
                recherche_status=RechercheStatus.NICHT_VERFUEGBAR
            ),
        ]

        # Synthese aus Stage 3
        self.sample_synthesis = RetailPriceSynthesis(
            geschaetzter_retail_preis_eur=285.0,
            preisspanne_min_eur=240.0,
            preisspanne_max_eur=350.0,
            median_web_preis_eur=300.0,
            anzahl_gefundene_preise=8,
            begruendung_preisfindung="Auktions- und eBay-Verkäufe zwischen 280 € und 310 € bilden das solide Preisniveau.",
            ausreisser_bereinigung_notiz="Pamono- und 1stDibs-Galeriepreise (750-850 €) als Händleraufschläge eingestuft.",
            produktbeschreibung="Verkaufsfertige Produktbeschreibung: Authentischer Vintage-Teaktisch aus den 1960er Jahren.",
            physische_merkmale={"material": "Teak massiv", "breite_cm": 80.0, "laenge_cm": 120.0, "hoehe_cm": 50.0},
            zustandsbericht={"zustand": "gut", "maengel": ["Kratzer 5cm"]}
        )

        # PipelineState erstellen
        self.state = PipelineState(
            item_name="Artikel_4092_Teaktisch",
            video_path=Path("input/Artikel_4092/video.mp4"),
            run_dir=self.test_dir,
            detected_id="ARTIKEL-4092",
            visual_analysis_json={
                "titel": "Vintage-Designertisch Mid-Century Teak",
                "kategorie": "Möbel",
                "produktbeschreibung": "Verkaufsfertige Produktbeschreibung: Authentischer Vintage-Teaktisch aus den 1960er Jahren.",
                "hersteller_oder_marke": "Dänisches Design",
                "modell_oder_epoche": "Mid-Century 1960er",
                "material": "Teak massiv",
                "farbe": "Teak Braun",
                "laenge_cm": 120.0,
                "breite_cm": 80.0,
                "hoehe_cm": 50.0,
                "maengel": ["Kratzer 5cm"]
            },
            reference_listings=[l.model_dump(mode="json") for l in self.sample_listings],
            discovered_web_sources=[l.model_dump(mode="json") for l in self.sample_listings],
            retail_price_synthesis_json=self.sample_synthesis.model_dump(mode="json"),
            web_price_summary=self.sample_synthesis.model_dump(mode="json")
        )

    def tearDown(self):
        shutil.rmtree(self.test_dir, ignore_errors=True)

    def test_build_web_research_rows_schema_and_hyperlink_formula(self):
        """Prüft, ob build_web_research_rows alle 9 geforderten Spalten und korrekte =HYPERLINK-Formeln erzeugt."""
        rows = ExportService.build_web_research_rows("ARTIKEL-4092", self.sample_listings)
        
        self.assertEqual(len(rows), 10)
        first_row = rows[0]
        
        expected_columns = [
            "Artikel_ID",
            "Website_Name",
            "Produkt_Titel",
            "Preis_EUR",
            "Preis_Typ",
            "Match_Genauigkeit",
            "Zustand_Referenz",
            "Status",
            "Link_URL"
        ]
        for col in expected_columns:
            self.assertIn(col, first_row, f"Fehlende Spalte in Web-Research-Zeile: {col}")

        self.assertEqual(first_row["Artikel_ID"], "ARTIKEL-4092")
        self.assertEqual(first_row["Website_Name"], "eBay")
        self.assertEqual(first_row["Preis_EUR"], 280.0)
        self.assertEqual(first_row["Preis_Typ"], PriceType.REALISIERTER_VERKAUFSPREIS.value)
        self.assertEqual(first_row["Match_Genauigkeit"], MatchGenauigkeit.EXAKTER_TREFFER.value)
        self.assertEqual(first_row["Status"], RechercheStatus.ERFOLGREICH.value)

        link_formula = first_row["Link_URL"]
        self.assertTrue(link_formula.startswith("=HYPERLINK("))
        self.assertIn("https://www.ebay.de/itm/123456789", link_formula)

        empty_url_row = rows[9]
        self.assertNotIn("=HYPERLINK", empty_url_row["Link_URL"])

    def test_build_unified_item_row_contains_valuation_summary_columns(self):
        """Prüft, ob build_unified_item_row die Bewertungsspalten auf Hauptempfehlung enthält."""
        row = ExportService.build_unified_item_row_from_state(self.state)

        self.assertEqual(row["Empfohlener_Retail_Preis_EUR"], 285.0)
        self.assertEqual(row["Preisspanne_Min_EUR"], 240.0)
        self.assertEqual(row["Preisspanne_Max_EUR"], 350.0)
        self.assertEqual(row["Median_Web_Preis_EUR"], 300.0)
        self.assertEqual(row["Anzahl_gefundene_Webpreise"], 8)
        self.assertIn("https://", row["Top_Referenz_Links"])
        self.assertIn("Auktions- und eBay-Verkäufe", row["Begruendung_Preisfindung"])

        self.assertIn("Verkaufsfertige Produktbeschreibung", row["produktbeschreibung"])
        self.assertEqual(row["material"], "Teak massiv")
        self.assertEqual(row["laenge_cm"], 120.0)
        self.assertEqual(row["breite_cm"], 80.0)
        self.assertEqual(row["hoehe_cm"], 50.0)
        self.assertEqual(row["maengel"], "Kratzer 5cm")

    def test_export_single_item_creates_2_sheets_with_web_prices(self):
        """Prüft, ob export_single_item genau die 2 Sheets Hauptempfehlung und Marktrecherche_Webpreise anlegt."""
        result = self.export_service.export_single_item(self.state, "test_item")
        excel_path = result["excel"]
        self.assertTrue(excel_path.exists())

        wb = openpyxl.load_workbook(excel_path)
        sheet_names = wb.sheetnames

        self.assertEqual(sheet_names, ["Hauptempfehlung", "Marktrecherche_Webpreise"])

        ws_web = wb["Marktrecherche_Webpreise"]
        self.assertEqual(ws_web.max_row, 11)  # 1 Header-Zeile + 10 Datenzeilen
        self.assertEqual(ws_web.max_column, 9)

        headers = [ws_web.cell(row=1, column=c).value for c in range(1, 10)]
        self.assertListEqual(headers, [
            "Artikel_ID",
            "Website_Name",
            "Produkt_Titel",
            "Preis_EUR",
            "Preis_Typ",
            "Match_Genauigkeit",
            "Zustand_Referenz",
            "Status",
            "Link_URL"
        ])

        link_cell = ws_web.cell(row=2, column=9)
        self.assert_valid_hyperlink_cell(link_cell, expected_url="https://www.ebay.de/itm/123456789")

    def assert_valid_hyperlink_cell(self, cell, expected_url: str = None):
        """Wiederverwendbare Hilfsmethode zur Validierung von =HYPERLINK-Formelzellen."""
        val = str(cell.value)
        self.assertTrue(val.startswith("=HYPERLINK("))
        if expected_url:
            self.assertIn(expected_url, val)
        self.assertEqual(cell.data_type, "f")
        self.assertEqual(cell.font.color.rgb, "FF0563C1")
        self.assertEqual(cell.font.underline, "single")

    def test_export_consolidated_batch_with_web_research(self):
        """Prüft, ob export_consolidated_batch Alle_Artikel und Marktrecherche_Webpreise sauber exportiert."""
        item_row = ExportService.build_unified_item_row_from_state(self.state)
        web_rows = ExportService.build_web_research_rows("ARTIKEL-4092", self.sample_listings)

        batch_result = self.export_service.export_consolidated_batch(
            items_rows=[item_row],
            web_research_rows=web_rows,
            base_filename="consolidated_test"
        )

        wb = openpyxl.load_workbook(batch_result["excel"])
        self.assertEqual(wb.sheetnames, ["Alle_Artikel", "Marktrecherche_Webpreise"])

        ws_web = wb["Marktrecherche_Webpreise"]
        self.assertEqual(ws_web.max_row, 11)
        link_cell = ws_web.cell(row=2, column=9)
        self.assert_valid_hyperlink_cell(link_cell, expected_url="https://www.ebay.de/itm/123456789")

    def test_empty_web_research_listings(self):
        """Prüft, ob bei leeren Referenz-Listings dennoch das Sheet mit Header-Spalten erzeugt wird."""
        self.state.reference_listings = []
        self.state.discovered_web_sources = []
        result = self.export_service.export_single_item(self.state, "test_empty_research")
        wb = openpyxl.load_workbook(result["excel"])
        self.assertIn("Marktrecherche_Webpreise", wb.sheetnames)
        ws_web = wb["Marktrecherche_Webpreise"]
        self.assertEqual(ws_web.max_row, 1)  # Nur Header
        self.assertEqual(ws_web.max_column, 9)

    def test_quotes_and_special_chars_in_hyperlink(self):
        """Prüft, ob Anführungszeichen im Titel in der =HYPERLINK Formel korrekt escaped werden."""
        listing = ReferenceListing(
            website_name="DesignStore",
            listing_titel='Vintage Teaktisch "Modell 42" Sonderedition',
            preis_eur=500.0,
            quell_url='https://example.com/item?id=42&cat="vintage"',
            recherche_status=RechercheStatus.ERFOLGREICH
        )
        rows = ExportService.build_web_research_rows("ART-1", [listing])
        formula = rows[0]["Link_URL"]
        self.assertTrue(formula.startswith("=HYPERLINK("))
        self.assertIn('""Modell 42""', formula)


if __name__ == "__main__":
    unittest.main()
