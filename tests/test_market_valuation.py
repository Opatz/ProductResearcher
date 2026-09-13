import json
import os
import sys
import tempfile
import shutil
import unittest
import subprocess
from pathlib import Path
from typing import Dict, Any, List
from unittest.mock import MagicMock

import openpyxl
import pandas as pd

from config.settings import AppConfig, GoogleSettings, PipelineSettings
from pipeline.models import (
    PriceType,
    MatchGenauigkeit,
    RechercheStatus,
    TargetWebsiteSuggestion,
    ReferenceListing,
    VisualAnalysisResult,
    RetailPriceSynthesis,
)
from pipeline.orchestrator import VideoLLMPipeline
from pipeline.state import PipelineState
from services.export_service import ExportService
from services.web_research_service import WebResearchService
from services.appraiser_service import AppraiserService


class TestPhase1And2Parsing(unittest.TestCase):
    """Prüft das Parsing von Phase 1 (Visuelle Analyse) und Phase 2 (10 Plattform-Recherchen)."""

    def setUp(self):
        self.web_service = WebResearchService(gemini_service=None)

    def test_phase1_visual_analysis_parsing_and_ten_targets(self):
        raw_payload = {
            "id": "ARTIKEL-100",
            "titel": "Kaiser Idell Schreibtischlampe Modell 6631",
            "kategorie": "Beleuchtung",
            "produktbeschreibung": "Originale Christian Dell Kaiser Idell Lampe in Schwarz mit Patina.",
            "hersteller_oder_marke": "Gebr. Kaiser & Co. / Christian Dell",
            "modell_oder_epoche": "Bauhaus 1930er",
            "physische_merkmale": {
                "hoehe_cm": 45.0,
                "durchmesser_cm": 28.5,
                "material": "Stahlblech lackiert, Messing"
            },
            "zustandsbericht": {
                "zustand": "gut",
                "maengel": ["Leichte Kratzer am Lampenschirm", "Originalverkabelung gealtert"]
            },
            "ziel_webseiten": [
                {
                    "website_name": "Pamono",
                    "target_url": "https://www.pamono.de/kaiser-idell",
                    "suchbegriff": "Kaiser Idell 6631 Christian Dell",
                    "plattform_typ": "design_galerie",
                    "begruendung": "Führende Plattform für Bauhaus-Designklassiker"
                },
                {
                    "website_name": "eBay",
                    "target_url": "https://www.ebay.de",
                    "suchbegriff": "Kaiser Idell 6631 Original",
                    "plattform_typ": "auktionshaus",
                    "begruendung": "Realisierte Verkäufe zur Marktpreisverifizierung"
                }
            ]
        }

        result = self.web_service.parse_visual_analysis_payload(raw_payload, ensure_ten=True)
        self.assertIsInstance(result, VisualAnalysisResult)
        self.assertEqual(result.titel, "Kaiser Idell Schreibtischlampe Modell 6631")
        self.assertEqual(result.kategorie, "Beleuchtung")
        self.assertEqual(result.hersteller_oder_marke, "Gebr. Kaiser & Co. / Christian Dell")
        self.assertEqual(result.physische_merkmale.get("hoehe_cm"), 45.0)
        self.assertEqual(len(result.zustandsbericht.get("maengel", [])), 2)
        
        self.assertEqual(len(result.ziel_webseiten), 10)
        site_names = [t.website_name for t in result.ziel_webseiten]
        self.assertIn("Pamono", site_names)
        self.assertIn("eBay", site_names)

    def test_phase2_site_research_response_parsing_success(self):
        target = TargetWebsiteSuggestion(
            website_name="Dorotheum",
            target_url="https://www.dorotheum.com/lot/9988",
            suchbegriff="Christian Dell Schreibtischleuchte"
        )
        object_summary = {
            "titel": "Kaiser Idell 6631",
            "hersteller_oder_marke": "Christian Dell",
            "modell_oder_epoche": "Bauhaus",
            "kategorie": "Beleuchtung",
            "geschaetzter_preis": 350.0
        }

        listing = self.web_service._create_mock_single_listing(
            site_name=target.website_name,
            target_url=target.target_url,
            object_summary=object_summary
        )

        self.assertIsInstance(listing, ReferenceListing)
        self.assertEqual(listing.website_name, "Dorotheum")
        self.assertIsNotNone(listing.preis_eur)
        self.assertGreater(listing.preis_eur, 0)
        self.assertEqual(listing.urspruengliche_waehrung, "EUR")
        self.assertEqual(listing.preis_typ, PriceType.REALISIERTER_VERKAUFSPREIS)
        self.assertEqual(listing.match_genauigkeit, MatchGenauigkeit.EXAKTER_TREFFER)
        self.assertEqual(listing.recherche_status, RechercheStatus.ERFOLGREICH)
        self.assertTrue(listing.quell_url.startswith("http"))

    def test_phase2_site_research_error_and_blocked_handling(self):
        target = TargetWebsiteSuggestion(
            website_name="BlockedPlatform",
            target_url="https://blocked.example.com",
            suchbegriff="Bauhaus Lampe"
        )
        object_summary = {"titel": "Test Lampe", "geschaetzter_preis": 200.0}

        class MockGeminiBlocked:
            def execute_text_prompt(self, *args, **kwargs):
                raise RuntimeError("403 Client Error: Forbidden for url: https://blocked.example.com")

        svc = WebResearchService(gemini_service=MockGeminiBlocked())
        listing = svc.research_single_site(target=target, object_summary=object_summary)

        self.assertEqual(listing.website_name, "BlockedPlatform")
        self.assertIsNone(listing.preis_eur)
        self.assertEqual(listing.recherche_status, RechercheStatus.ZUGRIFF_BLOCKIERT)
        self.assertEqual(listing.preis_typ, PriceType.UNBEKANNT)

    def test_phase2_parallel_mock_execution_returns_ten_listings(self):
        targets = [
            TargetWebsiteSuggestion(website_name=f"Site_{i}", target_url=f"https://site{i}.com")
            for i in range(1, 11)
        ]
        object_summary = {"titel": "Mid-Century Tisch", "geschaetzter_preis": 300.0}

        results = self.web_service.research_all_sites_parallel(
            targets=targets,
            object_summary=object_summary,
            max_workers=5,
            timeout=5.0
        )

        self.assertEqual(len(results), 10)
        for r in results:
            self.assertIsInstance(r, ReferenceListing)
            self.assertTrue(r.website_name.startswith("Site_"))


class TestAppraiserReconciliationAndDiscounting(unittest.TestCase):
    """Prüft die Gutachter-Schlichtung: Realisierte Preise vs. Angebotspreise & Mängelabschläge."""

    def setUp(self):
        self.appraiser_service = AppraiserService(gemini_service=None)

    def test_realized_price_priority_over_dealer_asking_prices(self):
        catalog = {
            "id": "ITEM-1",
            "titel": "Vintage Teak Esstisch",
            "kategorie": "Möbel",
            "geschaetzter_preis": 300.0,
            "zustandsbericht": {"zustand": "gut", "maengel": []}
        }
        listings = [
            ReferenceListing(
                website_name="eBay",
                listing_titel="Teak Esstisch - Verkauft",
                preis_eur=280.0,
                preis_typ=PriceType.REALISIERTER_VERKAUFSPREIS,
                match_genauigkeit=MatchGenauigkeit.EXAKTER_TREFFER,
                recherche_status=RechercheStatus.ERFOLGREICH
            ),
            ReferenceListing(
                website_name="Dorotheum",
                listing_titel="Teak Esstisch Hammerpreis",
                preis_eur=300.0,
                preis_typ=PriceType.REALISIERTER_VERKAUFSPREIS,
                match_genauigkeit=MatchGenauigkeit.EXAKTER_TREFFER,
                recherche_status=RechercheStatus.ERFOLGREICH
            ),
            ReferenceListing(
                website_name="Pamono",
                listing_titel="Dänischer Teaktisch Galeriepreis",
                preis_eur=900.0,
                preis_typ=PriceType.ANGEBOTSPREIS,
                match_genauigkeit=MatchGenauigkeit.EXAKTER_TREFFER,
                recherche_status=RechercheStatus.ERFOLGREICH
            ),
            ReferenceListing(
                website_name="1stDibs",
                listing_titel="Mid-Century Teak Tisch",
                preis_eur=1100.0,
                preis_typ=PriceType.ANGEBOTSPREIS,
                match_genauigkeit=MatchGenauigkeit.EXAKTER_TREFFER,
                recherche_status=RechercheStatus.ERFOLGREICH
            )
        ]

        synthesis = self.appraiser_service.synthesize_valuation(catalog, listings)
        self.assertIsInstance(synthesis, RetailPriceSynthesis)

        self.assertLess(synthesis.geschaetzter_retail_preis_eur, 500.0)
        self.assertGreater(synthesis.geschaetzter_retail_preis_eur, 200.0)
        self.assertIn("Händler", synthesis.ausreisser_bereinigung_notiz)

    def test_calibrated_condition_discounting(self):
        base_catalog = {
            "id": "ITEM-2",
            "titel": "Meissen Porzellanfigur",
            "kategorie": "Porzellan",
            "geschaetzter_preis": 500.0
        }
        intact_listings = [
            ReferenceListing(
                website_name="Auktionshaus",
                listing_titel="Meissen Figur Makellos",
                preis_eur=500.0,
                preis_typ=PriceType.REALISIERTER_VERKAUFSPREIS,
                match_genauigkeit=MatchGenauigkeit.EXAKTER_TREFFER,
                recherche_status=RechercheStatus.ERFOLGREICH
            )
        ]

        cat_damaged = dict(base_catalog)
        cat_damaged["zustandsbericht"] = {
            "zustand": "beschädigt",
            "maengel": ["Chip an der Hand", "Haarriss am Sockel"]
        }
        synth_damaged = self.appraiser_service.synthesize_valuation(cat_damaged, intact_listings)

        cat_scratched = dict(base_catalog)
        cat_scratched["zustandsbericht"] = {
            "zustand": "gebraucht",
            "maengel": ["Leichte Kratzer auf der Unterseite"]
        }
        synth_scratched = self.appraiser_service.synthesize_valuation(cat_scratched, intact_listings)

        cat_mint = dict(base_catalog)
        cat_mint["zustandsbericht"] = {
            "zustand": "neuwertig",
            "maengel": []
        }
        synth_mint = self.appraiser_service.synthesize_valuation(cat_mint, intact_listings)

        self.assertLess(synth_damaged.geschaetzter_retail_preis_eur, synth_scratched.geschaetzter_retail_preis_eur)
        self.assertLess(synth_scratched.geschaetzter_retail_preis_eur, synth_mint.geschaetzter_retail_preis_eur)
        self.assertIn("Mängelabschlag", synth_damaged.begruendung_preisfindung)


class TestMultiSheetExcelExportIntegration(unittest.TestCase):
    """Prüft den Export mit Marktrecherche_Webpreise, HYPERLINK-Formeln und Hauptempfehlung."""

    def setUp(self):
        self.test_dir = Path(tempfile.mkdtemp())
        self.export_service = ExportService(self.test_dir)

    def tearDown(self):
        shutil.rmtree(self.test_dir, ignore_errors=True)

    def test_multisheet_excel_contains_all_required_sheets_and_hyperlinks(self):
        state = PipelineState(
            item_name="Artikel_42",
            video_path=Path("dummy_42.mp4"),
            run_dir=self.test_dir
        )
        state.detected_id = "ARTIKEL-42"
        state.visual_analysis_json = {
            "titel": "Art Déco Tischuhr Junghans",
            "kategorie": "Uhren",
            "produktbeschreibung": "Feine Art Déco Kaminuhr aus Nussbaum.",
            "laenge_cm": 25.0,
            "breite_cm": 12.0,
            "hoehe_cm": 18.0,
            "material": "Nussbaum massiv, Messing",
            "maengel": ["Glas leicht berieben"]
        }
        state.reference_listings = [
            ReferenceListing(
                website_name="eBay",
                listing_titel="Junghans Art Déco Kaminuhr",
                preis_eur=195.0,
                urspruengliche_waehrung="EUR",
                preis_typ=PriceType.REALISIERTER_VERKAUFSPREIS,
                match_genauigkeit=MatchGenauigkeit.EXAKTER_TREFFER,
                zustand_referenz="Funktionsfähig mit Patina",
                quell_url="https://www.ebay.de/itm/55667788",
                recherche_status=RechercheStatus.ERFOLGREICH
            ),
            ReferenceListing(
                website_name="Pamono",
                listing_titel="Junghans Tischuhr 1930er",
                preis_eur=480.0,
                urspruengliche_waehrung="EUR",
                preis_typ=PriceType.ANGEBOTSPREIS,
                match_genauigkeit=MatchGenauigkeit.MODELLVARIANTE,
                zustand_referenz="Restauriert",
                quell_url="https://www.pamono.de/junghans-clock",
                recherche_status=RechercheStatus.ERFOLGREICH
            )
        ]
        state.retail_price_synthesis_json = {
            "geschaetzter_retail_preis_eur": 210.0,
            "preisspanne_min_eur": 175.0,
            "preisspanne_max_eur": 250.0,
            "median_web_preis_eur": 195.0,
            "anzahl_gefundene_preise": 2,
            "begruendung_preisfindung": "Guter Marktwert gestützt durch eBay-Verkauf.",
            "ausreisser_bereinigung_notiz": "Pamono Galeriepreis als spekulativer Aufschlag bewertet."
        }

        export_paths = self.export_service.export_single_item(state, base_filename="artikel_42_result")
        excel_path = export_paths["excel"]

        self.assertTrue(excel_path.exists())
        wb = openpyxl.load_workbook(excel_path)

        self.assertEqual(wb.sheetnames, ["Hauptempfehlung", "Marktrecherche_Webpreise"])

        ws_main = wb["Hauptempfehlung"]
        headers_main = [cell.value for cell in ws_main[1]]
        self.assertIn("Empfohlener_Retail_Preis_EUR", headers_main)
        self.assertIn("Preisspanne_Min_EUR", headers_main)
        self.assertIn("Preisspanne_Max_EUR", headers_main)
        self.assertIn("Median_Web_Preis_EUR", headers_main)
        self.assertIn("Anzahl_gefundene_Webpreise", headers_main)
        self.assertIn("Top_Referenz_Links", headers_main)
        self.assertIn("Begruendung_Preisfindung", headers_main)
        self.assertIn("produktbeschreibung", headers_main)
        self.assertIn("material", headers_main)

        ws_web = wb["Marktrecherche_Webpreise"]
        headers_web = [cell.value for cell in ws_web[1]]
        expected_web_headers = [
            "Artikel_ID", "Website_Name", "Produkt_Titel", "Preis_EUR",
            "Preis_Typ", "Match_Genauigkeit", "Zustand_Referenz", "Status", "Link_URL"
        ]
        for eh in expected_web_headers:
            self.assertIn(eh, headers_web)

        link_col_idx = headers_web.index("Link_URL") + 1
        formula_cell_val = ws_web.cell(row=2, column=link_col_idx).value
        self.assertTrue(str(formula_cell_val).startswith("=HYPERLINK("))
        self.assertIn("https://www.ebay.de/itm/55667788", str(formula_cell_val))


class TestEndToEndPipelineExecution(unittest.TestCase):
    """Prüft die vollständige Pipeline-Ausführung mit gemockten Services."""

    def setUp(self):
        self.test_root = Path(tempfile.mkdtemp())
        self.input_dir = self.test_root / "input"
        self.output_dir = self.test_root / "output"
        self.artikel_dir = self.input_dir / "artikel"
        self.artikel_dir.mkdir(parents=True, exist_ok=True)
        self.output_dir.mkdir(parents=True, exist_ok=True)

        self.item_folder = self.artikel_dir / "Artikel_01"
        self.item_folder.mkdir(parents=True, exist_ok=True)
        
        self.video_file = self.item_folder / "video.mp4"
        self.video_file.write_bytes(b"DUMMY_MP4_CONTENT")
        
        self.img1 = self.item_folder / "img_01.jpg"
        self.img1.write_bytes(b"DUMMY_JPG_1")
        self.img2 = self.item_folder / "img_02.jpg"
        self.img2.write_bytes(b"DUMMY_JPG_2")

        self.config = AppConfig(
            google=GoogleSettings(api_key="mock-key", model_name="gemini-3.5-flash-lite", enable_google_search=False),
            pipeline=PipelineSettings(
                raw_dir=self.input_dir / "raw",
                processed_dir=self.input_dir / "processed",
                artikel_dir=self.artikel_dir,
                input_dir=self.input_dir,
                context_dir=self.test_root / "context",
                output_dir=self.output_dir,
                preview_duration_sec=2.0,
                transcription_language="de"
            ),
            base_dir=Path(__file__).resolve().parent.parent
        )

    def tearDown(self):
        shutil.rmtree(self.test_root, ignore_errors=True)

    def test_pipeline_run_end_to_end(self):
        pipeline = VideoLLMPipeline(
            config=self.config,
            execution_dir=self.output_dir / "test_exec"
        )
        pipeline.video_service.extract_audio = MagicMock(return_value=self.output_dir / "audio.mp3")
        pipeline.video_service.trim_video = MagicMock(return_value=self.output_dir / "preview.mp4")

        mock_vis = VisualAnalysisResult(
            id="01",
            titel="Vintage Teaktisch",
            kategorie="Möbel",
            produktbeschreibung="Schöner Mid-Century Teaktisch",
            hersteller_oder_marke="Dänemark",
            modell_oder_epoche="1960er",
            geschaetztes_jahr_oder_epoche="ca. 1965",
            physische_merkmale={"material": "Teak massiv", "breite_cm": 80.0},
            zustandsbericht={"zustand": "gut", "maengel": []},
            ziel_webseiten=[TargetWebsiteSuggestion(website_name=f"Site_{i}", target_url=f"https://site{i}.com") for i in range(1, 11)]
        )
        pipeline.web_research_service.analyze_visual_and_suggest_targets = MagicMock(return_value=mock_vis)

        mock_listings = [
            ReferenceListing(
                website_name=f"Site_{i}",
                listing_titel=f"Listing {i}",
                preis_eur=250.0 + i * 10,
                preis_typ=PriceType.REALISIERTER_VERKAUFSPREIS if i <= 3 else PriceType.ANGEBOTSPREIS,
                match_genauigkeit=MatchGenauigkeit.EXAKTER_TREFFER,
                recherche_status=RechercheStatus.ERFOLGREICH,
                quell_url=f"https://site{i}.com/item"
            )
            for i in range(1, 11)
        ]
        pipeline.web_research_service.research_all_sites_parallel = MagicMock(return_value=mock_listings)

        mock_synthesis = RetailPriceSynthesis(
            geschaetzter_retail_preis_eur=285.0,
            preisspanne_min_eur=240.0,
            preisspanne_max_eur=350.0,
            median_web_preis_eur=300.0,
            anzahl_gefundene_preise=10,
            begruendung_preisfindung="Solide Marktlage.",
            ausreisser_bereinigung_notiz="Keine Ausreißer.",
            produktbeschreibung="Schöner Mid-Century Teaktisch",
            physische_merkmale={"material": "Teak massiv"},
            zustandsbericht={"zustand": "gut"}
        )
        pipeline.appraiser_service.synthesize_valuation = MagicMock(return_value=mock_synthesis)

        state = pipeline.run(
            video_path=self.video_file,
            image_paths=[self.img1, self.img2],
            item_name="Artikel_01"
        )

        self.assertEqual(state.status, "SUCCESS")
        item_dir = self.output_dir / "test_exec" / "Artikel_01"
        self.assertTrue(item_dir.exists())

        # 1. Zwischenschritt-Artefakte
        self.assertTrue((item_dir / "06_initial_visual_analysis.json").exists())
        self.assertTrue((item_dir / "06_web_research_10_sites.json").exists())
        self.assertTrue((item_dir / "06_retail_price_synthesis.json").exists())
        self.assertTrue((item_dir / "pipeline_trace.json").exists())

        # 2. Exportierte Dateien
        csv_path = item_dir / "video_result.csv"
        excel_path = item_dir / "video_result.xlsx"
        self.assertTrue(csv_path.exists())
        self.assertTrue(excel_path.exists())

        # 3. Excel-Inhalte validieren
        wb = openpyxl.load_workbook(excel_path)
        self.assertEqual(wb.sheetnames, ["Hauptempfehlung", "Marktrecherche_Webpreise"])

        ws_web = wb["Marktrecherche_Webpreise"]
        self.assertEqual(ws_web.max_row, 11)

        ws_main = wb["Hauptempfehlung"]
        self.assertEqual(ws_main.max_row, 2)
        headers = [c.value for c in ws_main[1]]
        self.assertIn("Empfohlener_Retail_Preis_EUR", headers)
        self.assertIn("Median_Web_Preis_EUR", headers)


if __name__ == "__main__":
    unittest.main()
